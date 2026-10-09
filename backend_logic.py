# backend_logic.py (CORRECTED - GRAMMAGE VALIDATION IMPROVED)

import openpyxl
import re
from datetime import datetime
import excel2img
import os
import time
import shutil
import json
import tkinter as tk
from tkinter import filedialog, messagebox
from openpyxl.drawing.image import Image as OpenpyxlImage

# Local dependencies
import config_manager as cm

# --- CONSTANTS ---
IMAGENES_GUARDADAS_FILE = os.path.join(cm.APP_DIR, "imagenes_guardadas.json")

def _realizar_calculos_internos(hoja):
    """
    Centralizes all calculations based on the final business rules.
    Reads raw data from the sheet and returns a dictionary with the calculated results.
    """
    resultados_calculados = {}
    
    def _convertir_a_numero(valor):
        """Safely converts a value to float, returning 0.0 on error."""
        if valor is None: return 0.0
        try: return float(valor)
        except (ValueError, TypeError): return 0.0

    # --- CORRECTED RULE 1: Sample Weight ---
    # Sum of all weight components, except Mortality.
    suma_componentes = [
        _convertir_a_numero(hoja['G30'].value), # Good Shrimp
        _convertir_a_numero(hoja['G65'].value), # Juveniles
        _convertir_a_numero(hoja['G67'].value)  # Foreign Material
    ]
    # Sum all defects from G32 to G61 (including Mouth Fungus)
    for i in range(32, 62):
        suma_componentes.append(_convertir_a_numero(hoja[f'G{i}'].value))
    
    peso_de_la_muestra = sum(suma_componentes)
    resultados_calculados['peso_de_la_muestra'] = peso_de_la_muestra
    
    # --- CORRECTED RULE 2: General Average Weight ---
    # Sum of Good Shrimp (G30) + Standard Defects (G32 to G59). Does NOT include mortality.
    suma_base_defectos = sum(_convertir_a_numero(hoja[f'G{i}'].value) for i in range(32, 60))
    camaron_bueno_bruto = _convertir_a_numero(hoja['G30'].value)
    
    peso_promedio_general = suma_base_defectos + camaron_bueno_bruto
    resultados_calculados['peso_promedio_general'] = peso_promedio_general
    
    if peso_promedio_general == 0:
        return resultados_calculados # Avoid division by zero

    # --- Percentage Calculations ---
    # Calculated based on the 'general average weight'.
    for i in range(32, 62): # Includes Mouth Fungus
        celda_g = f'G{i}'
        valor_bruto = _convertir_a_numero(hoja[celda_g].value)
        porcentaje = (valor_bruto / peso_promedio_general) * 100
        celda_h = celda_g.replace('G', 'H')
        resultados_calculados[celda_h] = porcentaje
        
    # Mortality Percentage
    mortalidad_bruto = _convertir_a_numero(hoja['G73'].value)
    porcentaje_mortalidad = (mortalidad_bruto / peso_promedio_general) * 100
    resultados_calculados['H73'] = porcentaje_mortalidad
    
    # Calculate percentages for Juveniles and Foreign Material based on total sample weight
    if peso_de_la_muestra > 0:
        juveniles_bruto = _convertir_a_numero(hoja['G65'].value)
        porcentaje_juveniles = (juveniles_bruto / peso_de_la_muestra) * 100
        resultados_calculados['H65_calc'] = porcentaje_juveniles

        material_bruto = _convertir_a_numero(hoja['G67'].value)
        porcentaje_material = (material_bruto / peso_de_la_muestra) * 100
        resultados_calculados['H67_calc'] = porcentaje_material
    else:
        resultados_calculados['H65_calc'] = 0.0
        resultados_calculados['H67_calc'] = 0.0

    # Decomposed Shrimp Percentage (uses N77)
    if peso_de_la_muestra > 0:
        descompuesto_bruto = _convertir_a_numero(hoja['N77'].value)
        porcentaje_descompuesto = (descompuesto_bruto / peso_de_la_muestra) * 100
        resultados_calculados['N77_calc'] = porcentaje_descompuesto
    else:
        resultados_calculados['N77_calc'] = 0.0
    
    return resultados_calculados

def analizar_archivo_para_gui(ruta_archivo, muestra_letra):
    """
    Main function to analyze an Excel file for the GUI, extracting data,
    running validations, and preparing the message.
    """
    try:
        wb = openpyxl.load_workbook(ruta_archivo, data_only=True, keep_vba=True)
    except Exception as e:
        return {"success": False, "error": f"File is open or corrupted. Please close it and try again.\nError: {e}"}

    hoja_nombre = next((h.title for h in wb.worksheets if h.sheet_state == 'visible' and h.title.strip().upper().endswith(muestra_letra)), None)
    if not hoja_nombre:
        wb.close()
        return {"success": False, "error": f"Sample '{muestra_letra}' not found in the file."}
        
    hoja = wb[hoja_nombre]
    resultados_calculados = _realizar_calculos_internos(hoja)

    lote_raw = hoja["F6"].value
    lote = ''.join(c for c in str(lote_raw) if c.isdigit() or c == '-') if lote_raw else "NO_BATCH"
    
    alertas_locales = []
    
    nombre_archivo = os.path.basename(ruta_archivo)
    match_lote_archivo = re.search(r'(\d{7}-\d{2})', nombre_archivo)
    if match_lote_archivo:
        lote_del_archivo = match_lote_archivo.group(1)
        lote_de_hoja = str(hoja["F6"].value or "").strip()
        if lote_del_archivo not in lote_de_hoja:
             alertas_locales.append("ALERTA: El lote del nombre del archivo no coincide con el lote de la hoja Excel.")

    # --- INICIO DE LA MODIFICACIÓN ---
    proveedor = hoja["F11"].value or "NO_PROVIDER"
    piscina = hoja["F14"].value or "NO_POOL"
    
    # Se llama a la nueva validación de consistencia con el historial del día
    alertas_locales.extend(validar_consistencia_con_historial(lote, proveedor, piscina, muestra_letra))
    # --- FIN DE LA MODIFICACIÓN ---
    
    alertas_locales.extend(validar_datos_obligatorios(hoja, muestra_letra, lote_raw, hoja_nombre))
    
    alertas_globales = validar_consistencia_global(ruta_archivo)

    certificacion = "ASC" if str(hoja["H10"].value).strip().lower() == "x" else "CONVENCIONAL"
    sabor, _ = analizar_sabor(hoja["L73"].value)
    es_entero = hoja["Q6"].value is not None
    
    defectos = extraer_defectos(hoja, resultados_calculados)
    hay_variacion_gramaje = verificar_variacion_gramaje(hoja)

    aceptables, cargadas, verdes, amarillo, gris, rojiza = None, None, None, None, None, None
    if muestra_letra == "A":
        aceptables = hoja["C77"].value
        cargadas = hoja["E77"].value
        verdes = hoja["G77"].value
        amarillo = hoja["H77"].value
        gris = hoja["I77"].value
        rojiza = hoja["L77"].value

    wb.close()
    
    condicion_sabor_anormal = (muestra_letra == "A" and "100% NORMAL" not in sabor)

    texto = f"""*LIBERACIÓN EN RECEPCIÓN*\nLOTE: {lote}\nPROVEEDOR: {proveedor}\nPISCINA: {piscina}\nMUESTRA: {muestra_letra}\nCERTIFICACIÓN: {certificacion}"""
    
    if hoja["Q11"].value is not None and str(hoja["Q11"].value).strip() != "":
        texto += "\nPESCA VIVA"
    if not es_entero:
        texto += "\nLOTE PARA COLA DIRECTA"
    
    texto += "\n**OBSERVACIONES**"
    
    texto += procesar_sabor_en_mensaje(sabor)
    if hay_variacion_gramaje:
        texto += "\n🚩 PRESENTA VARIACIÓN DE GRAMAJE"
    for defecto in defectos:
        texto += f"\n{defecto}"
    porcentaje_juveniles = resultados_calculados.get('H65_calc', 0.0)
    if porcentaje_juveniles >= 3:
        texto += f"\n🚩 JUVENILES {round(porcentaje_juveniles, 2)}%"
    porcentaje_material = resultados_calculados.get('H67_calc', 0.0)
    if porcentaje_material >= 3:
        texto += f"\n🚩 MATERIAL EXTRAÑO {round(porcentaje_material, 2)}%"
    if muestra_letra == "A":
        if any(isinstance(val, (int, float)) and val > 0 for val in [aceptables, cargadas, verdes]):
            texto += f"\n**COCCIÓN EN RECEPCIÓN**"
            if isinstance(aceptables, (int, float)) and aceptables > 0: texto += f"\n✅ CABEZAS ACEPTABLES {int(round(aceptables * 100, 0))}%"
            if isinstance(cargadas, (int, float)) and cargadas > 0: texto += f"\n🚩 CABEZAS CARGADAS {int(round(cargadas * 100, 0))}%"
            if isinstance(verdes, (int, float)) and verdes > 0: texto += f"\n🚩 CABEZAS VERDES {int(round(verdes * 100, 0))}%"
        
        tonalidades = []
        if isinstance(amarillo, (int, float)) and amarillo > 0: tonalidades.append(f"{int(round(amarillo * 100, 0))}% AMARILLO")
        if isinstance(gris, (int, float)) and gris > 0: tonalidades.append(f"{int(round(gris * 100, 0))}% GRIS")
        if isinstance(rojiza, (int, float)) and rojiza > 0: tonalidades.append(f"{int(round(rojiza * 100, 0))}% ROJIZA")
        if tonalidades:
            texto += "\n**VARIACIÓN DE TONALIDAD**\n" + " - ".join(tonalidades)
    
    if condicion_sabor_anormal:
        texto += "\n*SABOR APROBADO POR JEFATURA DE CALIDAD*\nREMUESTREADO POR: "
    
    return {
        "success": True,
        "message": texto,
        "alerts": {
            "local": sorted(list(set(alertas_locales))), # For Pop-up
            "global": sorted(list(set(alertas_globales))) # For panel only
        },
        "data": {
            "lote": lote,
            "proveedor": proveedor, # Asegúrate de devolver estos datos si los necesitas en la GUI
            "piscina": piscina,     # para la función de guardado.
            "hoja_nombre": hoja_nombre,
            "muestra": muestra_letra,
            "ask_remuestreador": False
        }
    }

def guardar_foto_adicional_muestra_a(ruta_foto_original, lote):
    """Saves a copy of an additional photo named after the batch."""
    try:
        config = cm.cargar_configuracion()
        carpeta_liberaciones = config.get("carpeta_liberaciones")
        if not carpeta_liberaciones: return None, "Release photo folder not configured."
        os.makedirs(carpeta_liberaciones, exist_ok=True)
        _, extension = os.path.splitext(ruta_foto_original)
        nombre_archivo_destino = f"{lote}{extension}"
        ruta_destino = os.path.join(carpeta_liberaciones, nombre_archivo_destino)
        shutil.copy2(ruta_foto_original, ruta_destino)
        print(f"✅ Additional photo for Sample A saved as: {nombre_archivo_destino}")
        return ruta_destino, None
    except Exception as e: 
        print(f"⛔ Error saving additional photo: {e}")
        return None, str(e)

def tomar_captura_para_gui(ruta_archivo, hoja_nombre, lote):
    """Takes a screenshot of a specific range in Excel."""
    try:
        config = cm.cargar_configuracion()
        carpeta_fotos = config.get("carpeta_fotos")
        if not carpeta_fotos: return {"success": False, "error": "Temporary screenshot folder not configured."}
        os.makedirs(carpeta_fotos, exist_ok=True)
        nombre_imagen = f"{lote}_{hoja_nombre}_{int(time.time())}.png"
        ruta_imagen = os.path.join(carpeta_fotos, nombre_imagen)
        excel2img.export_img(ruta_archivo, ruta_imagen, hoja_nombre, "C2:R87")
        if os.path.exists(ruta_imagen) and os.path.getsize(ruta_imagen) > 5 * 1024:
            return {"success": True, "path": ruta_imagen}
        else:
            if os.path.exists(ruta_imagen): os.remove(ruta_imagen)
            return {"success": False, "error": "Generated image is empty or the process failed."}
    except Exception as e:
        return {"success": False, "error": f"Error taking screenshot: {e}"}

def obtener_muestras_disponibles(ruta_archivo):
    """Reads an Excel file and returns a list of available samples (A, B, C, etc.)."""
    if not ruta_archivo or not os.path.exists(ruta_archivo):
        return {"success": False, "error": "File does not exist."}
    muestras = []
    try:
        wb = openpyxl.load_workbook(ruta_archivo, read_only=True)
        for sheet_name in wb.sheetnames:
            match = re.search(r'MUESTRA\s+([A-Z])$', sheet_name.strip().upper())
            if match: muestras.append(match.group(1))
        wb.close()
        return {"success": True, "muestras": sorted(muestras)}
    except Exception as e:
        return {"success": False, "error": f"Could not read the Excel file.\n\nError: {e}"}

# --- Full Validation and Extraction Logic ---

def _get_base_lote(lote_string):
    """
    Helper function to extract the base batch number, removing the sample letter.
    Ex: '0001234-26 A' -> '0001234-26'
    """
    if not lote_string:
        return ""
    parts = lote_string.strip().rsplit(' ', 1)
    if len(parts) == 2 and len(parts[1]) == 1 and parts[1].isalpha():
        return parts[0]
    return lote_string.strip()

def validar_consistencia_global(ruta_archivo):
    """
    Compares master data, color profiles, and grammages across all sample sheets.
    """
    alertas_globales = []
    try:
        wb = openpyxl.load_workbook(ruta_archivo, data_only=True, keep_vba=True)
    except Exception as e:
        print(f"Could not open {ruta_archivo} for global validation: {e}")
        return []

    muestras = [s for s in wb.worksheets if s.sheet_state == 'visible' and re.search(r'MUESTRA\s+([A-Z])$', s.title.strip().upper())]
    
    if len(muestras) <= 1:
        wb.close()
        return []

    # 1. Master Data Validation
    try:
        hoja_ref = muestras[0]
        lote_ref_full = str(hoja_ref["F6"].value or "").strip()
        datos_ref = {
            'Lote': _get_base_lote(lote_ref_full),
            'Proveedor': str(hoja_ref["F11"].value or "").strip(),
            'Piscina': str(hoja_ref["F14"].value or "").strip(),
            'Tipo de Producto': 'Entero' if hoja_ref["Q6"].value is not None else 'Cola',
            'Certificación': 'ASC' if hoja_ref["H10"].value is not None else 'Convencional',
            'Pesca Viva': 'Sí' if hoja_ref["Q11"].value is not None else 'No',
            'nombre_hoja': hoja_ref.title
        }
        colores_ref = { 'claros': (float(hoja_ref['L65'].value or 0)) + (float(hoja_ref['M65'].value or 0)), 'oscuros': float(hoja_ref['O65'].value or 0) }
        datos_ref['Perfil Color'] = 'Claro' if colores_ref['claros'] > colores_ref['oscuros'] else 'Oscuro'

        for hoja_comp in muestras[1:]:
            lote_comp_full = str(hoja_comp["F6"].value or "").strip()
            datos_comp = {
                'Lote': _get_base_lote(lote_comp_full),
                'Proveedor': str(hoja_comp["F11"].value or "").strip(),
                'Piscina': str(hoja_comp["F14"].value or "").strip(),
                'Tipo de Producto': 'Entero' if hoja_comp["Q6"].value is not None else 'Cola',
                'Certificación': 'ASC' if hoja_comp["H10"].value is not None else 'Convencional',
                'Pesca Viva': 'Sí' if hoja_comp["Q11"].value is not None else 'No',
                'nombre_hoja': hoja_comp.title
            }
            colores_comp = { 'claros': (float(hoja_comp['L65'].value or 0)) + (float(hoja_comp['M65'].value or 0)), 'oscuros': float(hoja_comp['O65'].value or 0) }
            datos_comp['Perfil Color'] = 'Claro' if colores_comp['claros'] > colores_comp['oscuros'] else 'Oscuro'
            
            for campo, valor_ref in datos_ref.items():
                if campo == 'nombre_hoja': continue
                valor_comp = datos_comp.get(campo)
                if valor_ref != valor_comp:
                    display_ref, display_comp = (lote_ref_full, lote_comp_full) if campo == 'Lote' else (valor_ref, valor_comp)
                    alerta_msg = (f"ALERTA: El perfil de color es inconsistente entre Muestra {datos_ref['nombre_hoja'].split()[-1]} "
                                 f"(predomina {valor_ref.lower()}) y Muestra {datos_comp['nombre_hoja'].split()[-1]} (predomina {valor_comp.lower()})." 
                                 if campo == 'Perfil Color' else 
                                 f"ALERTA: Inconsistencia en '{campo}'. Muestra {datos_ref['nombre_hoja'].split()[-1]} es '{display_ref}' y Muestra {datos_comp['nombre_hoja'].split()[-1]} es '{display_comp}'.")
                    if alerta_msg not in alertas_globales:
                         alertas_globales.append(alerta_msg)
    except Exception as e:
        print(f"Error reading master data: {e}")
        alertas_globales.append(f"ALERT: Could not read data from sheet '{muestras[0].title}' for comparison.")

    # 2. UPDATED VALIDATION: Identical grammages (3 or more)
    celdas_gramaje = {"O31": "100-120","O32": "80-100","O33": "70-80","O34": "60-70","O35": "50-60","O36": "40-50","O37": "30-40","O38": "20-30","O39": "10-20"}
    for i in range(len(muestras)):
        for j in range(i + 1, len(muestras)):
            hoja1 = muestras[i]
            hoja2 = muestras[j]
            letra1 = hoja1.title.strip().split()[-1]
            letra2 = hoja2.title.strip().split()[-1]
            tallas_identicas = []
            
            for celda, talla in celdas_gramaje.items():
                try:
                    valor1 = float(hoja1[celda].value or 0)
                    valor2 = float(hoja2[celda].value or 0)
                    
                    if valor1 > 0 and valor1 == valor2:
                        tallas_identicas.append(talla)
                except (ValueError, TypeError):
                    continue # Ignore cells with text or incorrect format
            
            # The alert is triggered only if 3 or more sizes are identical
            if len(tallas_identicas) >= 3:
                alerta_gramaje = (f"ALERTA: Gramaje idéntico entre Muestra {letra1} y Muestra {letra2} en {len(tallas_identicas)} tallas: "
                                  f"{', '.join(tallas_identicas)}.")
                if alerta_gramaje not in alertas_globales:
                    alertas_globales.append(alerta_gramaje)

    wb.close()
    return alertas_globales


def analizar_sabor(celda_sabor_valor):
    """Analyzes the flavor cell value and validates that percentages sum to 100%."""
    if not celda_sabor_valor:
        return "SIN DATOS DE SABOR", "⚠️ FALTA INFORMACIÓN DE SABOR"
    
    texto_sabor = str(celda_sabor_valor).strip().upper()
    porcentajes = re.findall(r'(\d+(?:\.\d+)?)%', texto_sabor)
    
    if not porcentajes:
        return texto_sabor, None
    
    try:
        numeros = [float(p) for p in porcentajes]
        suma_total = sum(numeros)
        
        if abs(suma_total - 100.0) > 0.1: # Use a small tolerance
            alerta = f"⚠️ SABOR INVÁLIDO: Suma {suma_total}% - Debe ser exactamente 100%"
            return texto_sabor, alerta
        else:
            return texto_sabor, None
    except ValueError:
        return texto_sabor, "⚠️ ERROR EN FORMATO DE PORCENTAJES DE SABOR"

def validar_color(hoja):
    """Validates that the sum of the color cells is exactly 20."""
    celdas_color = ["L65", "M65", "N65", "O65", "P65", "Q65"]
    valores_color = []
    
    for celda in celdas_color:
        valor = hoja[celda].value
        if valor is not None:
            try:
                valores_color.append(float(valor))
            except (ValueError, TypeError):
                return f"⚠️ ERROR EN FORMATO DE COLOR: Celda {celda} contiene valor inválido"
    
    if not valores_color:
        return "⚠️ FALTA INFORMACIÓN DE COLOR: Ninguna celda de color tiene datos"
    
    suma_color = sum(valores_color)
    
    if suma_color != 20:
        return f"⚠️ COLOR INCORRECTO: Suma {suma_color} - Debe ser exactamente 20"
    
    return None

def validar_lote_vs_hoja(lote_valor, nombre_hoja):
    """Validates that the batch letter matches the sheet letter."""
    if not lote_valor: return "⚠️ FALTA INFORMACIÓN DEL LOTE"
    if not nombre_hoja: return "⚠️ NOMBRE DE HOJA INVÁLIDO"
    
    lote_str = str(lote_valor).strip().upper()
    match_lote = re.search(r'\s+([A-Z])$', lote_str) # Search for letter at the end
    if not match_lote: return "⚠️ FALTA INDICAR MUESTRA EN EL LOTE (ej: ...-25 A)"
    
    letra_lote = match_lote.group(1)
    
    nombre_hoja_upper = str(nombre_hoja).strip().upper()
    match_hoja = re.search(r'MUESTRA\s+([A-Z])$', nombre_hoja_upper)
    if not match_hoja: return "⚠️ NOMBRE DE HOJA INVÁLIDO (ej: MUESTRA A)"
    
    letra_hoja = match_hoja.group(1)
    
    if letra_lote != letra_hoja:
        return f"⚠️ INCONSISTENCIA: Lote es muestra '{letra_lote}', pero hoja es '{letra_hoja}'"
    
    return None

def validar_consistencia_con_historial(lote_actual, proveedor_actual, piscina_actual, muestra_actual):
    """
    Checks the current data against messages sent today to prevent inconsistencies.
    """
    alertas = []
    registro_mensajes = _cargar_mensajes_enviados()
    hoy = datetime.now().strftime('%Y-%m-%d')

    # Extraer el número base del lote actual (sin la letra de muestra) para comparaciones
    base_lote_actual = ''.join(c for c in str(lote_actual) if c.isdigit() or c == '-')

    for datos_enviados in registro_mensajes.values():
        # Solo comparar con los registros de hoy
        if datos_enviados.get('fecha') != hoy:
            continue

        lote_enviado = datos_enviados.get('lote', '')
        base_lote_enviado = ''.join(c for c in str(lote_enviado) if c.isdigit() or c == '-')
        proveedor_enviado = datos_enviados.get('proveedor', '')
        piscina_enviada = datos_enviados.get('piscina', '')

        # CASO 1: Misma piscina, diferente lote (Solo para Muestra A)
        # "ayer libere... el lote 0001240-26... piscina 20 muestra a y luego envie... 0001241-26... piscina 20 muestra a"
        if muestra_actual == 'A' and piscina_actual == piscina_enviada and base_lote_actual != base_lote_enviado:
            alerta = (f"ALERTA: Ya se envió el lote '{lote_enviado}' de la piscina '{piscina_actual}' hoy. "
                      f"¿Es correcto liberar un nuevo lote de la misma piscina?")
            if alerta not in alertas:
                alertas.append(alerta)

        # Comprobar solo si los lotes base son iguales para los siguientes casos
        if base_lote_actual == base_lote_enviado:
            # CASO 2: Mismo lote, diferente piscina
            # "o el mismo lote con diferente piscina informe"
            if piscina_actual != piscina_enviada:
                alerta = (f"ALERTA: El lote '{lote_actual}' ya se envió hoy desde la piscina '{piscina_enviada}'. "
                          f"¿Confirma que ahora corresponde a la piscina '{piscina_actual}'?")
                if alerta not in alertas:
                    alertas.append(alerta)
            
            # CASO 3: Mismo lote, diferente proveedor
            # "o el mismo lote con diferente proveedor informe"
            if proveedor_actual != proveedor_enviado:
                alerta = (f"ALERTA: El lote '{lote_actual}' ya se envió hoy con el proveedor '{proveedor_enviado}'. "
                          f"¿Confirma que ahora corresponde al proveedor '{proveedor_actual}'?")
                if alerta not in alertas:
                    alertas.append(alerta)
                    
    return alertas


def validar_datos_obligatorios(hoja, muestra_letra, lote_valor=None, nombre_hoja=None):
    """Executes all data validations on the sheet and returns a list of alerts."""
    alertas = []

    # --- NEW INDIVIDUAL VALIDATIONS ---
    # Mortality vs. Live Catch Validation
    es_pesca_viva = hoja['Q11'].value is not None
    valor_mortalidad = hoja['G73'].value
    if valor_mortalidad is not None and isinstance(valor_mortalidad, (int, float)) and valor_mortalidad > 0:
        if not es_pesca_viva:
            alertas.append("ALERTA: Se ingresaron datos de mortalidad sin ser Pesca Viva.")

    # "Good Shrimp" Validation
    camaron_bueno = hoja['G30'].value
    if camaron_bueno is None or (isinstance(camaron_bueno, (int, float)) and camaron_bueno <= 0):
        alertas.append('ALERTA: Falta el peso del "Camarón Bueno" (celda G30), dato indispensable.')
        
    # Color Incompatibility Validation
    color_a1_val = hoja['L65'].value or 0
    color_a2_val = hoja['M65'].value or 0
    color_a4_val = hoja['O65'].value or 0
    
    tiene_claros = (isinstance(color_a1_val, (int, float)) and color_a1_val > 0) or \
                   (isinstance(color_a2_val, (int, float)) and color_a2_val > 0)
    tiene_oscuros = isinstance(color_a4_val, (int, float)) and color_a4_val > 0
    
    if tiene_claros and tiene_oscuros:
        alertas.append("ALERTA: Se detectó color oscuro (A4) junto a colores claros (A1/A2), lo cual es improbable.")
    
    # --- EXISTING VALIDATIONS ---
    if not hoja["F11"].value or str(hoja["F11"].value).strip() == "": alertas.append("⚠️ FALTA INFORMACIÓN DEL PROVEEDOR")
    if not hoja["F14"].value or str(hoja["F14"].value).strip() == "": alertas.append("⚠️ FALTA INFORMACIÓN DE LA PISCINA")
    if not hoja["G14"].value or (isinstance(hoja["G14"].value, (int, float)) and hoja["G14"].value <= 0): alertas.append("⚠️ FALTA INFORMACIÓN DE LIBRAS")
    
    celdas_horas = ["G79", "G81", "G83", "O79", "O81", "O83"]
    horas_faltantes = [c for c in celdas_horas if hoja[c].value is None or str(hoja[c].value).strip() == ""]
    if horas_faltantes: alertas.append(f"⚠️ FALTAN HORAS EN CELDAS: {', '.join(horas_faltantes)}")
    
    celdas_gramaje = {"O31": "100-120","O32": "80-100","O33": "70-80","O34": "60-70","O35": "50-60","O36": "40-50","O37": "30-40","O38": "20-30","O39": "10-20"}
    rangos_gramaje = {"100-120": (9, 10.99),"80-100": (11, 12.99),"70-80": (13, 14.99),"60-70": (15, 16.99),"50-60": (17, 19.99),"40-50": (20, 24.99),"30-40": (25, 33.99),"20-30": (34, 49.99),"10-20": (50, 100)}
    
    total_libras_gramaje = sum(hoja[c].value for c, t in celdas_gramaje.items() if isinstance(hoja[c].value, (int, float)) and hoja[c].value > 0)
    
    if total_libras_gramaje > 0:
        try:
            talla_predominante = max(((t, hoja[c].value) for c, t in celdas_gramaje.items() if isinstance(hoja[c].value, (int, float))), key=lambda item: item[1])[0]
            gramaje_promedio = (total_libras_gramaje * 454) / 100
            talla_calculada = next((t for t, (min_g, max_g) in rangos_gramaje.items() if min_g <= gramaje_promedio <= max_g), min(rangos_gramaje.keys(), key=lambda k: abs(gramaje_promedio - sum(rangos_gramaje[k])/2)))
            if talla_calculada != talla_predominante:
                alertas.append(f"⚠️ GRAMAJE NO CONCLUYENTE: Calculado ~{gramaje_promedio:.2f}g (talla {talla_calculada}), pero predomina talla {talla_predominante}")
        except ValueError:
             alertas.append("⚠️ No se pudo determinar la talla predominante del gramaje.")

    if hoja["Q6"].value is None and hoja["Q10"].value is None: alertas.append("⚠️ FALTA INDICAR SI ES ENTERO O COLA")
    if hoja["F10"].value is None and hoja["H10"].value is None: alertas.append("⚠️ FALTA INDICAR SI ES CERTIFICACIÓN CONVENCIONAL O ASC")
    if hoja["H87"].value is None and hoja["I87"].value is None: alertas.append("⚠️ FALTA INFORMACIÓN DE PLEÓPODOS")

    _, alerta_sabor = analizar_sabor(hoja["L73"].value)
    if alerta_sabor: alertas.append(alerta_sabor)

    alerta_color = validar_color(hoja)
    if alerta_color: alertas.append(alerta_color)
    
    if lote_valor and nombre_hoja:
        alerta_lote_hoja = validar_lote_vs_hoja(lote_valor, nombre_hoja)
        if alerta_lote_hoja: alertas.append(alerta_lote_hoja)

    if muestra_letra == "A":
        es_entero = hoja["Q6"].value is not None
        if es_entero:
            tiene_datos_coccion = any([hoja["C77"].value is not None, hoja["E77"].value is not None, hoja["G77"].value is not None])
            if not tiene_datos_coccion:
                alertas.append("⚠️ FALTA INFORMACIÓN DE COCCIÓN EN RECEPCIÓN (ACEPTABLES, CARGADAS O VERDES)")
            else:
                aceptables = hoja["C77"].value if isinstance(hoja["C77"].value, (int, float)) else 0
                cargadas = hoja["E77"].value if isinstance(hoja["E77"].value, (int, float)) else 0
                verdes = hoja["G77"].value if isinstance(hoja["G77"].value, (int, float)) else 0
                suma_coccion = aceptables + cargadas + verdes
                if abs(suma_coccion - 1.0) > 0.01:
                    porcentaje_suma = round(suma_coccion * 100, 1)
                    if suma_coccion < 1.0:
                        alertas.append(f"⚠️ COCCIÓN INCOMPLETA: Solo suma {porcentaje_suma}% - Debe ser 100%")
                    else:
                        alertas.append(f"⚠️ COCCIÓN EXCEDE: Suma {porcentaje_suma}% - Debe ser 100%")
        
        tiene_datos_tonalidad = any([hoja["H77"].value is not None, hoja["I77"].value is not None, hoja["L77"].value is not None])
        if not tiene_datos_tonalidad:
             alertas.append("⚠️ FALTA INDICAR LA TONALIDAD")
        else:
            amarillo = hoja["H77"].value if isinstance(hoja["H77"].value, (int, float)) else 0
            gris = hoja["I77"].value if isinstance(hoja["I77"].value, (int, float)) else 0
            rojiza = hoja["L77"].value if isinstance(hoja["L77"].value, (int, float)) else 0
            suma_tonalidad = amarillo + gris + rojiza
            if abs(suma_tonalidad - 1.0) > 0.01:
                porcentaje_tonalidad = round(suma_tonalidad * 100, 1)
                if suma_tonalidad < 1.0:
                    alertas.append(f"⚠️ TONALIDAD INCOMPLETA: Solo suma {porcentaje_tonalidad}% - Debe ser 100%")
                else:
                    alertas.append(f"⚠️ TONALIDAD EXCEDE: Suma {porcentaje_tonalidad}% - Debe ser 100%")

    return alertas

def extraer_defectos(hoja, resultados_calculados):
    """Extracts all defects that exceed defined thresholds, using calculated percentages."""
    celdas_defectos = {"H32":"FLACIDEZ","H33":"MUDADO","H35":"MANCHAS LEVES","H36":"MANCHAS FUERTES","H48":"CABEZAS NARANJAS","H50":"CABEZA DESCOLGADA","H51":"HEPAT PANCREAS REGADO","H52":"HEPAT PANCREAS REVENTADO","H54":"BRANQUIAS OSCURAS LEVES","H55":"BRANQUIAS AMARILLAS LEVES","H57":"BRANQUIAS OSCURAS FUERTES","H58":"BRANQUIAS AMARILLAS FUERTES","H59":"BRANQUIAS SUCIAS","H60":"HONGO BUCAL LEVE","H61":"HONGO BUCAL FUERTE"}
    defectos = []
    
    for celda_h, nombre in celdas_defectos.items():
        valor_calculado = resultados_calculados.get(celda_h, 0.0)
        if valor_calculado > 10:
            defectos.append(f"🚩 {nombre} {round(valor_calculado, 2)}%")
            
    valor_mortalidad = resultados_calculados.get('H73', 0.0)
    if valor_mortalidad > 0:
        defectos.append(f"🚩 MORTALIDAD {round(valor_mortalidad, 2)}%")
        
    valor_descompuesto = resultados_calculados.get('N77_calc', 0.0)
    if valor_descompuesto > 0:
        defectos.append(f"🚩 CAMARÓN EN DESCOMPOSICIÓN {round(valor_descompuesto, 2)}%")
        
    if isinstance(hoja["H87"].value, (int, float)) and hoja["H87"].value > 0:
        porcentaje_final = int(round(hoja['H87'].value * 100, 0)) if 0 < hoja['H87'].value <= 1 else int(round(hoja['H87'].value, 0))
        if porcentaje_final > 0:
            defectos.append(f"🚩 PLEOPODOS ROJOS {porcentaje_final}%")
            
    return defectos

def verificar_variacion_gramaje(hoja):
    """Checks if there are 5 or more grammage cells with values."""
    celdas_gramaje = [f"O{i}" for i in range(31, 40)]
    return sum(1 for c in celdas_gramaje if hoja[c].value is not None and str(hoja[c].value).strip() != "") >= 5

def procesar_sabor_en_mensaje(sabor_texto):
    """Formats the flavor text for the final message."""
    es_normal = "100% NORMAL" in sabor_texto
    return "\n✅ SABOR 100% NORMAL" if es_normal else f"\n🚩 SABOR {sabor_texto}"

# --- Excel Image Management Logic (Tab 2) ---

def insertar_imagen_openpyxl(ruta_excel, ruta_imagen, celda_ancla='AN18'):
    """Inserts an image into a specific cell of an Excel file."""
    try:
        wb = openpyxl.load_workbook(ruta_excel, keep_vba=True)
        ws = next((wb[sheet_name] for sheet_name in wb.sheetnames if sheet_name.strip() == 'MUESTRA A'), wb.active)
        img = OpenpyxlImage(ruta_imagen)
        img.width, img.height = 1083, 767
        img.anchor = celda_ancla
        ws.add_image(img)
        wb.save(ruta_excel)
        print(f"✅ Image inserted into '{os.path.basename(ruta_excel)}'.")
        return True
    except Exception as e:
        print(f"❌ Error inserting image: {e}")
        return False

def cargar_imagenes_guardadas():
    """Loads the record of saved images from a JSON file."""
    try:
        if os.path.exists(IMAGENES_GUARDADAS_FILE):
            with open(IMAGENES_GUARDADAS_FILE, 'r', encoding='utf-8') as f: return json.load(f)
        return {}
    except: return {}

def guardar_registro_imagen(lote, ruta_excel, ruta_imagen, estado="guardada"):
    """Saves a record of a processed image to the JSON file."""
    try:
        registro = cargar_imagenes_guardadas()
        timestamp = int(time.time())
        id_registro = f"{lote}_{os.path.basename(ruta_excel)}_{timestamp}"
        registro[id_registro] = { "lote": lote, "excel": os.path.basename(ruta_excel), "imagen": os.path.basename(ruta_imagen), "estado": estado, "fecha": datetime.now().isoformat(), "timestamp": timestamp}
        if len(registro) > 200:
            registros_ordenados = sorted(registro.items(), key=lambda x: x[1]['timestamp'], reverse=True)
            registro = dict(registros_ordenados[:200])
        with open(IMAGENES_GUARDADAS_FILE, 'w', encoding='utf-8') as f:
            json.dump(registro, f, indent=2)
    except Exception as e:
        print(f"Error saving record: {e}")

def verificar_imagen_ya_guardada(lote, ruta_excel):
    """Checks if an image for a batch/excel has already been saved today."""
    registro = cargar_imagenes_guardadas()
    hoy = datetime.now().strftime('%Y-%m-%d')
    for reg_id, datos in registro.items():
        if (datos['lote'] == lote and datos['excel'] == os.path.basename(ruta_excel) and datos['fecha'].startswith(hoy) and datos['estado'] == "guardada"):
            return True, reg_id
    return False, None

def buscar_correspondencias_imagenes():
    """Searches for Excel files and finds corresponding images by batch number."""
    config = cm.cargar_configuracion()
    carpeta_excel = config["carpeta_excel"]
    carpeta_fotos = config["carpeta_liberaciones"]
    if not all([carpeta_excel, carpeta_fotos, os.path.exists(carpeta_excel), os.path.exists(carpeta_fotos)]):
        print("❌ ERROR: Excel or photo folders not configured or not found.")
        return []
    
    archivos_excel = [f for f in os.listdir(carpeta_excel) if f.endswith((".xlsx", ".xlsm")) and not f.startswith("~")]
    imagenes_disponibles = {os.path.splitext(f)[0]: {'ruta': os.path.join(carpeta_fotos, f), 'archivo': f} for f in os.listdir(carpeta_fotos) if f.lower().endswith(('.png', '.jpg', '.jpeg'))}
    
    if not archivos_excel or not imagenes_disponibles: return []
    
    correspondencias = []
    for archivo_excel in archivos_excel:
        ruta_excel = os.path.join(carpeta_excel, archivo_excel)
        try:
            wb = openpyxl.load_workbook(ruta_excel, data_only=True)
            hoja_a = next((h for h in wb.worksheets if h.sheet_state == 'visible' and h.title.strip().upper().endswith('A')), None)
            if hoja_a and hoja_a["F6"].value:
                lote = ''.join(c for c in str(hoja_a["F6"].value) if c.isdigit() or c == '-')
                ya_guardada, reg_id = verificar_imagen_ya_guardada(lote, ruta_excel)
                imagen_encontrada = next((info for nombre, info in imagenes_disponibles.items() if lote in nombre or nombre in lote), None)
                correspondencias.append({'excel': ruta_excel, 'lote': lote, 'imagen': imagen_encontrada, 'archivo_excel': archivo_excel, 'ya_guardada': ya_guardada})
            wb.close()
        except Exception as e:
            print(f"⚠️ Error reading {archivo_excel}: {e}")
            
    return correspondencias

def procesar_imagenes_automatico():
    """Automatically processes all pending images."""
    print("--- 📸 AUTOMATIC IMAGE SAVING ---")
    correspondencias = buscar_correspondencias_imagenes()
    if not correspondencias:
        print("⚠️ No files or images found to process.")
        return
    pendientes = [c for c in correspondencias if not c['ya_guardada'] and c['imagen']]
    if not pendientes:
        print("\n🎉 All images are up to date!")
        return
    
    print(f"\n⏳ Found {len(pendientes)} pending images to process.")
    if messagebox.askyesno("Confirm Process", f"Do you want to process {len(pendientes)} pending images?"):
        exitosos, errores = 0, 0
        for i, corr in enumerate(pendientes, 1):
            print(f"\n--- Processing {i}/{len(pendientes)}: {corr['lote']} ---")
            if insertar_imagen_openpyxl(corr['excel'], corr['imagen']['ruta']):
                exitosos += 1
                guardar_registro_imagen(corr['lote'], corr['excel'], corr['imagen']['ruta'], "guardada")
            else:
                errores += 1
                guardar_registro_imagen(corr['lote'], corr['excel'], corr['imagen']['ruta'], "error")
            time.sleep(1)
        print(f"\n--- SUMMARY --- \n✅ Success: {exitosos}\n❌ Errors: {errores}")
        messagebox.showinfo("Process Finished", f"Process completed.\n\nSuccess: {exitosos}\nErrors: {errores}")

def procesar_una_imagen_manual():
    """Allows the user to manually select an Excel file and an image to process."""
    print("--- 🎯 MANUAL IMAGE SAVING ---")
    config = cm.cargar_configuracion()
    carpeta_excel = config["carpeta_excel"]
    if not carpeta_excel or not os.path.exists(carpeta_excel):
        messagebox.showerror("Error", "Excel folder not configured.")
        return
        
    ruta_excel = filedialog.askopenfilename(title="Select the Excel file", initialdir=carpeta_excel, filetypes=[("Excel files", "*.xlsx *.xlsm")])
    if not ruta_excel: return
    
    ruta_imagen = filedialog.askopenfilename(title="Select the image to insert", initialdir=config.get("carpeta_liberaciones"), filetypes=[("Image files", "*.png *.jpg *.jpeg")])
    if not ruta_imagen: return
    
    lote = "MANUAL_BATCH"
    try:
        wb = openpyxl.load_workbook(ruta_excel, data_only=True)
        hoja_a = next((h for h in wb.worksheets if h.sheet_state == 'visible' and h.title.strip().upper().endswith('A')), None)
        if hoja_a and hoja_a["F6"].value:
            lote = ''.join(c for c in str(hoja_a["F6"].value) if c.isdigit() or c == '-')
        wb.close()
    except Exception as e:
        print(f"Could not read the batch from excel: {e}")
        
    if messagebox.askyesno("Confirm Insertion", f"Insert the image:\n{os.path.basename(ruta_imagen)}\n\nInto the file:\n{os.path.basename(ruta_excel)}?"):
        if insertar_imagen_openpyxl(ruta_excel, ruta_imagen):
            guardar_registro_imagen(lote, ruta_excel, ruta_imagen, "guardada")
            messagebox.showinfo("Success", "Image inserted correctly.")
        else:
            guardar_registro_imagen(lote, ruta_excel, ruta_imagen, "error")
            messagebox.showerror("Error", "Could not insert the image.")

def mostrar_registro_imagenes():
    """Shows the history of saved images, grouped by date."""
    print("--- 📋 SAVED IMAGES LOG ---")
    registro = cargar_imagenes_guardadas()
    if not registro:
        print("📝 No images registered.")
        return
        
    registros_por_fecha = {}
    for datos in registro.values():
        fecha = datos['fecha'][:10]
        registros_por_fecha.setdefault(fecha, []).append(datos)
        
    for fecha in sorted(registros_por_fecha.keys(), reverse=True):
        print(f"\n📅 {fecha}:")
        for datos in sorted(registros_por_fecha[fecha], key=lambda x: x['timestamp'], reverse=True):
            estado = "✅" if datos['estado'] == "guardada" else "❌"
            print(f"   {estado} {datos['lote']} -> {datos['excel']} -> {datos['imagen']}")

def mostrar_correspondencias():
    """Shows an analysis of which images correspond to which Excel files without processing them."""
    print("--- 🔍 CORRESPONDENCE ANALYSIS ---")
    correspondencias = buscar_correspondencias_imagenes()
    if not correspondencias:
        print("📝 No files found to analyze.")
        return
        
    pendientes = [c for c in correspondencias if not c['ya_guardada'] and c['imagen']]
    sin_imagen = [c for c in correspondencias if not c['imagen']]
    
    if pendientes:
        print(f"\n⏳ PENDING PROCESSING ({len(pendientes)}):")
        for corr in pendientes:
            print(f"   - Batch {corr['lote']} in '{corr['archivo_excel']}' -> Matches '{corr['imagen']['archivo']}'")
    else:
        print("\n✅ No pending images to process.")
        
    if sin_imagen:
        print(f"\n❌ NO CORRESPONDING IMAGE ({len(sin_imagen)}):")
        for corr in sin_imagen:
            print(f"   - Batch {corr['lote']} in '{corr['archivo_excel']}' -> No image found.")

def limpiar_registro_imagenes():
    """Clears the entire record of saved images."""
    print("--- 🗑️ CLEAR LOG ---")
    if messagebox.askyesno("Confirm Clear", "Are you sure you want to delete the ENTIRE saved images log? This action cannot be undone."):
        try:
            if os.path.exists(IMAGENES_GUARDADAS_FILE):
                os.remove(IMAGENES_GUARDADAS_FILE)
            print("✅ Log cleared successfully.")
            messagebox.showinfo("Success", "The images log has been cleared.")
        except Exception as e:
            print(f"❌ Error clearing: {e}")
            messagebox.showerror("Error", f"Could not clear the log:\n{e}")

MENSAJES_ENVIADOS_FILE = os.path.join(cm.APP_DIR, "mensajes_enviados.json")

def _cargar_mensajes_enviados():
    """Loads the record of sent messages from a JSON file."""
    try:
        if os.path.exists(MENSAJES_ENVIADOS_FILE):
            with open(MENSAJES_ENVIADOS_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
    except (IOError, json.JSONDecodeError):
        return {}
    return {}

def _guardar_registro_mensajes(registro):
    """Saves the message log to the JSON file."""
    try:
        if len(registro) > 200:
            registros_ordenados = sorted(registro.items(), key=lambda x: x[1]['timestamp'], reverse=True)
            registro = dict(registros_ordenados[:200])
            
        with open(MENSAJES_ENVIADOS_FILE, 'w', encoding='utf-8') as f:
            json.dump(registro, f, indent=2)
    except IOError as e:
        print(f"Error saving message log: {e}")

def verificar_mensaje_duplicado(lote, muestra):
    """
    Checks if a message for a specific batch and sample has already been sent today.
    """
    registro = _cargar_mensajes_enviados()
    hoy = datetime.now().strftime('%Y-%m-%d')
    id_registro = f"{lote}_{muestra}_{hoy}"
    
    return id_registro in registro

def guardar_mensaje_enviado(lote, muestra, proveedor, piscina):
    """
    Saves a record that a message has been sent for a batch and sample,
    including provider and pool for consistency checks.
    """
    registro = _cargar_mensajes_enviados()
    hoy = datetime.now().strftime('%Y-%m-%d')
    timestamp = int(time.time())
    # The key remains the same for easy duplicate checking on the same day
    id_registro = f"{lote}_{muestra}_{hoy}"
    
    registro[id_registro] = {
        "lote": lote,
        "muestra": muestra,
        "proveedor": proveedor, # <-- DATO ADICIONAL
        "piscina": piscina,     # <-- DATO ADICIONAL
        "fecha": hoy,
        "timestamp": timestamp
    }
    _guardar_registro_mensajes(registro)
    print(f"✅ Message for Batch {lote} - Sample {muestra} (P: {proveedor}, L: {piscina}) registered as sent.")